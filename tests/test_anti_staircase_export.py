import pytest
import os
import openpyxl
from app.database.db_manager import DatabaseManager
from app.database.repository import Repository
from app.database.models import PageType, ExportConfig
from app.exporters.xlsx_exporter import XLSXExporter
from app.exporters.csv_exporter import CSVExporter

@pytest.fixture
def voter_repo(tmp_path):
    db_file = str(tmp_path / "voters.db")
    mgr = DatabaseManager(db_file)
    mgr.initialize()
    repo = Repository(mgr)

    doc = repo.add_document("bazidpur_pacs_voter_list.pdf", "/tmp/voter.pdf", 3, PageType.SCANNED, 10240)
    
    # 6 voter records across pages with identical canonical fields + stray OCR noise
    voters = [
        {"serial_no": "1", "voter_id": "ABC1234567", "name": "राम कुमार", "relation_name": "मोहन लाल", "house_no": "12", "age": "45", "gender": "पुरुष", "6": "sets5"},
        {"serial_no": "2", "voter_id": "ABC1234568", "name": "श्याम सिंह", "relation_name": "सोहन सिंह", "house_no": "14", "age": "32", "gender": "पुरुष", "yas": "-विद्याधरनगर"},
        {"serial_no": "3", "voter_id": "XYZ9876543", "name": "गीता देवी", "relation_name": "राम कुमार", "house_no": "12", "age": "40", "gender": "महिला", "ae": ",तकरी"},
        {"serial_no": "4", "voter_id": "XYZ9876544", "name": "राजेश कुमार", "relation_name": "सुरेश कुमार", "house_no": "18", "age": "29", "gender": "पुरुष", "ee": "/"},
        {"serial_no": "5", "voter_id": "XYZ9876545", "name": "सुनीता देवी", "relation_name": "राजेश कुमार", "house_no": "18", "age": "26", "gender": "महिला", "679": "बाज़िदपुर"},
        {"serial_no": "6", "voter_id": "XYZ9876546", "name": "अमित सिंह", "relation_name": "विमलेश सिंह", "house_no": "20", "age": "35", "gender": "पुरुष", "929": "ज्ञात"},
    ]

    for idx, v in enumerate(voters):
        page_num = (idx // 2) + 1
        repo.add_record(doc.id, page_num, None, v, f"Voter {idx+1}", 0, 0, 0, 0)

    repo.rebuild_search_index(doc.id)
    return repo, doc

def test_xlsx_no_diagonal_staircase(voter_repo, tmp_path):
    repo, doc = voter_repo
    out_xlsx = str(tmp_path / "voter_export.xlsx")
    exporter = XLSXExporter(repo)
    config = ExportConfig(format="xlsx", scope="all", category_id=None, record_ids=None, search_query=None, output_path=out_xlsx)
    exporter.export(config)

    assert os.path.exists(out_xlsx)
    wb = openpyxl.load_workbook(out_xlsx)
    ws = wb["All Records"]

    # Verify headers in Row 1
    headers = [cell.value for cell in ws[1] if cell.value is not None]
    
    # Check that stray OCR noise keys are NOT present as columns
    for bad_key in ['6', 'yas', 'ae', 'ee', '679', '929']:
        assert bad_key not in headers

    # Verify canonical headers are present in correct order
    assert headers[0] == "Serial No (क्र. सं.)"
    assert headers[1] == "Voter ID / EPIC"
    assert headers[2] == "Name (नाम)"
    assert headers[3] == "Father / Relative Name (पिता/सम्बन्धी का नाम)"
    assert headers[4] == "House No / Village (मकान/ग्राम)"
    assert headers[5] == "Age (उम्र)"
    assert headers[6] == "Gender (लिंग)"

    # Verify data in Rows 2 to 7 are ALIGNED (not staggered diagonally!)
    for row_idx in range(2, 8):
        sl = ws.cell(row=row_idx, column=1).value
        epic = ws.cell(row=row_idx, column=2).value
        name = ws.cell(row=row_idx, column=3).value
        rel = ws.cell(row=row_idx, column=4).value
        house = ws.cell(row=row_idx, column=5).value
        age = ws.cell(row=row_idx, column=6).value
        gender = ws.cell(row=row_idx, column=7).value

        # Every row must have its core fields populated in columns 1..7
        assert sl is not None and str(sl) == str(row_idx - 1)
        assert epic is not None and len(str(epic)) == 10
        assert name is not None and len(str(name)) > 2
        assert rel is not None and len(str(rel)) > 2
        assert house is not None
        assert age is not None
        assert gender in ["पुरुष", "महिला"]

def test_csv_aligned_export(voter_repo, tmp_path):
    repo, doc = voter_repo
    out_csv = str(tmp_path / "voter_export.csv")
    exporter = CSVExporter(repo)
    config = ExportConfig(format="csv", scope="all", category_id=None, record_ids=None, search_query=None, output_path=out_csv)
    exporter.export(config)

    assert os.path.exists(out_csv)
    with open(out_csv, 'r', encoding='utf-8-sig') as f:
        lines = [line.strip() for line in f if line.strip()]

    header_line = lines[0]
    assert "Serial No" in header_line
    assert "Voter ID" in header_line
    assert "Name" in header_line
    assert "679" not in header_line

    # Verify each voter data row contains voter info
    assert "राम कुमार" in lines[1]
    assert "श्याम सिंह" in lines[2]
    assert "गीता देवी" in lines[3]
