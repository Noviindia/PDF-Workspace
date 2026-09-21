import os
import zipfile
import html
from typing import List, Dict
from app.exporters.base_exporter import BaseExporter
from app.database.models import Record, Category

try:
    from docx import Document as DocxDocument
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
except (ImportError, Exception):
    DocxDocument = None


class DOCXExporter(BaseExporter):
    """Word document export using python-docx with pure-Python OpenXML fallback."""
    
    def export_records(self, records: List[Record], output_path: str, document_name: str, categories: Dict[int, Category]):
        if DocxDocument is not None:
            try:
                self._export_with_python_docx(records, output_path, document_name, categories)
                return
            except Exception as e:
                print(f"python-docx export failed ({e}), using OpenXML fallback...")
        
        self._export_pure_openxml(records, output_path, document_name, categories)

    def _export_with_python_docx(self, records: List[Record], output_path: str, document_name: str, categories: Dict[int, Category]):
        doc = DocxDocument()
        
        title = doc.add_heading(document_name, 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        fields = self._get_field_names(records)
        headers = fields + ['Category', 'Source Page']
        
        category_records = {}
        for record in records:
            cat_id = record.category_id
            if cat_id not in category_records:
                category_records[cat_id] = []
            category_records[cat_id].append(record)
            
        if not category_records:
            doc.add_paragraph("No records found.")
        
        for cat_id, cat_recs in category_records.items():
            cat = categories.get(cat_id)
            cat_name = cat.name if cat else "Uncategorized"
            
            doc.add_heading(cat_name, level=2)
            doc.add_paragraph(f"Records: {len(cat_recs)}", style='Subtitle')
            
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = 'Table Grid'
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            
            hdr_cells = table.rows[0].cells
            for i, header in enumerate(headers):
                hdr_cells[i].text = str(header)
                for paragraph in hdr_cells[i].paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = True
            
            for record in cat_recs:
                row_cells = table.add_row().cells
                values = self._get_record_values(record, fields)
                values.append(cat_name)
                values.append(str(record.source_page) if record.source_page else "")
                
                for i, value in enumerate(values):
                    row_cells[i].text = str(value)
                    
            doc.add_paragraph()
            
        doc.add_page_break()
        doc.add_heading('Summary', level=1)
        doc.add_paragraph(f"Total Records: {len(records)}")
        doc.add_paragraph(f"Total Categories: {len(category_records)}")
        
        pages = set(r.source_page for r in records if r.source_page)
        doc.add_paragraph(f"Total Source Pages: {len(pages)}")
        
        doc.save(output_path)

    def _export_pure_openxml(self, records: List[Record], output_path: str, document_name: str, categories: Dict[int, Category]):
        """Pure-Python OpenXML DOCX generator requiring zero external dependencies."""
        fields = self._get_field_names(records)
        headers = fields + ['Category', 'Source Page']

        category_records = {}
        for record in records:
            cat_id = record.category_id
            if cat_id not in category_records:
                category_records[cat_id] = []
            category_records[cat_id].append(record)

        body_xml = []
        body_xml.append(f'<w:p><w:pPr><w:jc w:val="center"/></w:pPr><w:r><w:rPr><w:b/><w:sz w:val="48"/></w:rPr><w:t>{html.escape(document_name)}</w:t></w:r></w:p>')

        if not category_records:
            body_xml.append('<w:p><w:r><w:t>No records found.</w:t></w:r></w:p>')

        for cat_id, cat_recs in category_records.items():
            cat = categories.get(cat_id)
            cat_name = cat.name if cat else "Uncategorized"

            body_xml.append(f'<w:p><w:r><w:rPr><w:b/><w:sz w:val="32"/></w:rPr><w:t>{html.escape(cat_name)}</w:t></w:r></w:p>')
            body_xml.append(f'<w:p><w:r><w:rPr><w:i/><w:sz w:val="22"/><w:color w:val="666666"/></w:rPr><w:t>Records: {len(cat_recs)}</w:t></w:r></w:p>')

            tbl = ['<w:tbl>', '<w:tblPr><w:tblBorders><w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/><w:left w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/><w:right w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/><w:insideV w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/></w:tblBorders></w:tblPr>']
            
            # Header row
            tbl.append('<w:tr><w:trPr><w:tblHeader/></w:trPr>')
            for h in headers:
                tbl.append(f'<w:tc><w:tcPr><w:shd w:val="clear" w:color="auto" w:fill="2E75B6"/></w:tcPr><w:p><w:r><w:rPr><w:b/><w:color w:val="FFFFFF"/></w:rPr><w:t>{html.escape(str(h))}</w:t></w:r></w:p></w:tc>')
            tbl.append('</w:tr>')

            # Data rows
            for record in cat_recs:
                tbl.append('<w:tr>')
                values = self._get_record_values(record, fields)
                values.append(cat_name)
                values.append(str(record.source_page) if record.source_page else "")
                for val in values:
                    tbl.append(f'<w:tc><w:p><w:r><w:t>{html.escape(str(val))}</w:t></w:r></w:p></w:tc>')
                tbl.append('</w:tr>')

            tbl.append('</w:tbl>')
            body_xml.append(''.join(tbl))
            body_xml.append('<w:p/>')

        # Summary
        body_xml.append('<w:p><w:r><w:br w:type="page"/></w:r></w:p>')
        body_xml.append('<w:p><w:r><w:rPr><w:b/><w:sz w:val="32"/></w:rPr><w:t>Summary</w:t></w:r></w:p>')
        body_xml.append(f'<w:p><w:r><w:t>Total Records: {len(records)}</w:t></w:r></w:p>')
        body_xml.append(f'<w:p><w:r><w:t>Total Categories: {len(category_records)}</w:t></w:r></w:p>')
        pages = set(r.source_page for r in records if r.source_page)
        body_xml.append(f'<w:p><w:r><w:t>Total Source Pages: {len(pages)}</w:t></w:r></w:p>')

        document_xml = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    {''.join(body_xml)}
  </w:body>
</w:document>'''

        content_types_xml = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>'''

        rels_xml = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>'''

        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        with zipfile.ZipFile(output_path, 'w', compression=zipfile.ZIP_DEFLATED) as zf:
            zf.writestr('[Content_Types].xml', content_types_xml)
            zf.writestr('_rels/.rels', rels_xml)
            zf.writestr('word/document.xml', document_xml)

