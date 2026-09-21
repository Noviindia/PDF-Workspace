import os
from typing import List, Dict
try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None

from app.exporters.base_exporter import BaseExporter
from app.database.models import Record, Category

class XLSXExporter(BaseExporter):
    """Excel export using openpyxl."""
    
    def export_records(self, records: List[Record], output_path: str, document_name: str, categories: Dict[int, Category]):
        if not openpyxl:
            raise ImportError("openpyxl is required for XLSX export.")
            
        wb = openpyxl.Workbook()
        ws_all = wb.active
        ws_all.title = "All Records"
        
        self._write_sheet(ws_all, records, "All Records", categories)
        
        # Group by category
        category_records = {}
        for record in records:
            cat_id = record.category_id
            if cat_id not in category_records:
                category_records[cat_id] = []
            category_records[cat_id].append(record)
            
        if len(category_records) > 0 and len(records) > 0:
            for cat_id, cat_recs in category_records.items():
                cat = categories.get(cat_id)
                cat_name = cat.name if cat else "Uncategorized"
                safe_cat_name = self._sanitize_filename(cat_name)[:31] # Excel sheet name limit
                if not safe_cat_name:
                    safe_cat_name = f"Category_{cat_id}"
                
                # Check for duplicate names
                base_name = safe_cat_name
                count = 1
                while safe_cat_name in wb.sheetnames:
                    safe_cat_name = f"{base_name[:28]}_{count}"
                    count += 1
                    
                ws_cat = wb.create_sheet(title=safe_cat_name)
                self._write_sheet(ws_cat, cat_recs, f"{document_name} - {cat_name}", categories)
                
        wb.save(output_path)
        
    def _write_sheet(self, ws, records: List[Record], title: str, categories: Dict[int, Category]):
        fields = self._get_field_names(records)
        display_headers = [self._get_display_field_name(f) for f in fields] + ['Category', 'Source Page']
        
        # Write headers
        for col_idx, header in enumerate(display_headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 28
        
        # Write data
        for row_idx, record in enumerate(records, 2):
            values = self._get_record_values(record, fields)
            cat = categories.get(record.category_id)
            cat_name = cat.name if cat else ""
            values.append(cat_name)
            page_val = record.page_number if hasattr(record, 'page_number') and record.page_number else (getattr(record, 'source_page', '') or '')
            values.append(page_val)
            
            for col_idx, value in enumerate(values, 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                
                # Try to convert to number
                try:
                    if value and str(value).replace('.', '', 1).isdigit():
                        if '.' in str(value):
                            cell.value = float(value)
                        else:
                            cell.value = int(value)
                    else:
                        cell.value = value
                except ValueError:
                    cell.value = value
                    
        # Auto-size columns
        for col_idx in range(1, len(display_headers) + 1):
            max_len = 0
            col_letter = get_column_letter(col_idx)
            for cell in ws[col_letter]:
                try:
                    if len(str(cell.value)) > max_len:
                        max_len = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_len + 2, 50)
            ws.column_dimensions[col_letter].width = adjusted_width
            
        # Add auto-filter
        if len(records) > 0:
            ws.auto_filter.ref = ws.dimensions
